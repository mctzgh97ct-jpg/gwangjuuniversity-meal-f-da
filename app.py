from flask import Flask, render_template, jsonify
import io
import re
from datetime import date, datetime
from urllib.parse import parse_qs, urljoin, urlparse, urlunparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

app = Flask(__name__)

BOARD_URL = "https://www.gwangju.ac.kr/bbs/?b_id=gwangju_jinwol_rm&mn=553&site=gwangju"
BOARD_AJAX_URL = "https://www.gwangju.ac.kr/bbs/bbs_ajax/?b_id=gwangju_jinwol_rm&mn=553&site=gwangju"
BASE_URL = "https://www.gwangju.ac.kr"
BOARD_ID = "gwangju_jinwol_rm"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; GwangjuMealVibe/1.0)"
}


def to_ajax_url(url):
    """게시판 껍데기 페이지(/bbs/)를 실제 목록·본문이 있는 /bbs/bbs_ajax/로 바꿉니다."""
    parsed = urlparse(url)
    path = parsed.path.rstrip("/") or "/"
    if path == "/bbs":
        return urlunparse(parsed._replace(path="/bbs/bbs_ajax/"))
    return url


def get_latest_post():
    """광주대학교 식당메뉴 게시판에서 가장 최근 식단 게시글을 찾습니다."""
    r = requests.get(BOARD_AJAX_URL, headers=HEADERS, timeout=15)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")

    candidates = []
    for a in soup.find_all("a", href=True):
        full = urljoin(BASE_URL, a["href"])
        query = parse_qs(urlparse(full).query)
        b_id = (query.get("b_id") or [""])[0]
        post_type = (query.get("type") or [""])[0]
        bs_idx = (query.get("bs_idx") or [""])[0]
        if b_id != BOARD_ID or post_type != "view" or not bs_idx.isdigit():
            continue
        title = " ".join(a.stripped_strings)
        candidates.append((int(bs_idx), title, full))

    if not candidates:
        raise RuntimeError("식단 게시글을 찾지 못했습니다. 학교 홈페이지 구조가 바뀌었을 수 있습니다.")

    candidates.sort(key=lambda item: item[0], reverse=True)
    _, title, url = candidates[0]
    return {"title": title, "url": url}


def find_excel_attachment(post_url):
    """게시글에서 xlsx/xls 첨부파일 다운로드 링크를 찾습니다."""
    r = requests.get(to_ajax_url(post_url), headers=HEADERS, timeout=15)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")

    possible = []
    for a in soup.find_all("a", href=True):
        text = " ".join(a.stripped_strings)
        href = a["href"]
        joined = urljoin(BASE_URL, href)
        haystack = f"{text} {href}".lower()

        if ".xlsx" in haystack or ".xls" in haystack:
            possible.append(joined)
        elif any(k in haystack for k in ["download", "file_down", "filedown", "attach"]):
            if "파일" in text or "xlsx" in text.lower() or "xls" in text.lower():
                possible.append(joined)

    if not possible:
        raise RuntimeError("엑셀 첨부파일 링크를 찾지 못했습니다.")

    return possible[0]


def download_workbook(url):
    r = requests.get(url, headers=HEADERS, timeout=20)
    r.raise_for_status()
    return r.content


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", "\n").split())


WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def worksheet_grid(ws):
    """병합 셀은 시작 셀 값으로 채워 실제 보이는 표와 같게 만듭니다."""
    grid = []
    for r in range(1, ws.max_row + 1):
        grid.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
    for merged in ws.merged_cells.ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                grid[r - 1][c - 1] = value
    return grid


def find_date_columns(grid):
    best_row = None
    best_cols = []
    for r_idx, row in enumerate(grid):
        cols = []
        for c_idx, value in enumerate(row):
            parsed = as_date(value)
            if parsed:
                cols.append((c_idx, parsed))
        if len(cols) > len(best_cols):
            best_row = r_idx
            best_cols = cols
    if best_row is None or not best_cols:
        raise RuntimeError("엑셀에서 식단 날짜 행을 찾지 못했습니다.")
    return best_row, best_cols


def is_set_menu_label(value):
    return clean(value) in {"정식", "학생정식"}


def pick_default_date(days, today):
    today_s = today.isoformat()
    with_menu = [day for day in days if day["items"]]
    candidates = with_menu or days
    if not candidates:
        return None
    for day in candidates:
        if day["date"] == today_s:
            return today_s
    for day in candidates:
        if day["date"] >= today_s:
            return day["date"]
    return candidates[-1]["date"]


def parse_student_meals(binary, today=None):
    """현재 첨부 엑셀에서 날짜열과 '정식' 구간을 찾아 학생정식만 뽑습니다."""
    if today is None:
        today = date.today()

    wb = load_workbook(io.BytesIO(binary), data_only=True)
    ws = wb[wb.sheetnames[0]]
    grid = worksheet_grid(ws)
    date_row_idx, date_cols = find_date_columns(grid)
    date_col_indexes = {c_idx for c_idx, _ in date_cols}

    set_rows = []
    for r_idx, row in enumerate(grid):
        if r_idx <= date_row_idx:
            continue
        labels = [row[c_idx] for c_idx in range(len(row)) if c_idx not in date_col_indexes]
        if any(is_set_menu_label(value) for value in labels):
            set_rows.append(r_idx)

    if not set_rows:
        raise RuntimeError("엑셀에서 학생정식 구간을 찾지 못했습니다.")

    days = []
    for c_idx, day in date_cols:
        items = []
        seen = set()
        for r_idx in set_rows:
            raw = grid[r_idx][c_idx] if c_idx < len(grid[r_idx]) else None
            if as_date(raw):
                continue
            text = clean(raw)
            if not text or is_set_menu_label(text) or text.lower() == "lunch":
                continue
            if text in seen:
                continue
            seen.add(text)
            items.append(text)
        days.append({
            "date": day.isoformat(),
            "weekday": WEEKDAYS[day.weekday()],
            "items": items,
        })

    return {
        "days": days,
        "default_date": pick_default_date(days, today),
    }


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/latest")
def api_latest():
    try:
        post = get_latest_post()
        attachment = find_excel_attachment(post["url"])
        binary = download_workbook(attachment)
        meals = parse_student_meals(binary)

        return jsonify({
            "ok": True,
            "source": BOARD_URL,
            "post": post,
            "attachment": attachment,
            "days": meals["days"],
            "default_date": meals["default_date"],
        })
    except Exception as e:
        return jsonify({
            "ok": False,
            "error": str(e),
            "source": BOARD_URL
        }), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)
